"""Academic table-and-figure layout following the user's monochrome reference."""
from pathlib import Path
import sys
import subprocess
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
TMP = ROOT/'tmp/bfp-survey/black-white'
TMP.mkdir(parents=True,exist_ok=True)
OUT = ROOT/'tests/bfp-mock-survey-black-and-white.docx'
RUNTIME = Path('C:/Users/janna/.cache/codex-runtimes/codex-primary-runtime/dependencies')
sys.path.insert(0,'C:/Users/janna/.codex/plugins/cache/openai-primary-runtime/documents/26.727.11326/skills/documents/scripts')
from table_geometry import apply_table_geometry

resident = [27,25,26,24,28,27,29,26,28,25]
bfp = [9,8,9,10,9,8,9,10,9,9]
yes,no = sum(resident)+sum(bfp), 400-sum(resident)-sum(bfp)
assert (yes,no)==(355,45)

# A monochrome pie, with an outline and labels to keep both categories distinct.
c=canvas.Canvas(str(TMP/'pie.pdf'),pagesize=(420,335))
c.setFillGray(0)
c.setFont('Times-Bold',14)
c.drawCentredString(210,314,'Grand Overall Survey Results')
c.setFont('Times-Roman',10)
c.drawCentredString(210,297,'Residents and BFP Personnel (Mock Data)')
c.setLineWidth(.8)
c.setStrokeGray(0)
c.setFillGray(.88)
c.wedge(96,48,330,282,startAng=90,extent=yes/400*360,stroke=1,fill=1)
c.setFillGray(.3)
c.wedge(96,48,330,282,startAng=90-no/400*360,extent=no/400*360,stroke=1,fill=1)
c.setFillGray(0)
c.setLineWidth(.6)
c.lines([(256,261,290,280),(290,280,341,280),(151,72,120,41),(120,41,82,41)])
c.setFont('Times-Bold',11)
c.drawString(345,278,'NO')
c.drawString(345,264,'11.25%')
c.drawRightString(79,41,'YES')
c.drawRightString(79,27,'88.75%')
c.setFont('Times-Roman',9)
c.drawCentredString(210,8,'355 Yes responses and 45 No responses | 400 item responses')
c.save()
subprocess.run([str(RUNTIME/'native/poppler/Library/bin/pdftoppm.exe'),'-png','-gray','-singlefile','-scale-to','2100',str(TMP/'pie.pdf'),str(TMP/'pie')],check=True,creationflags=0x08000000)

doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Inches(8.5),Inches(11)
sec.top_margin,sec.bottom_margin=Inches(1.2),Inches(1)
sec.left_margin=sec.right_margin=Inches(1.25)
sec.header_distance,sec.footer_distance=Inches(.4),Inches(.45)
sec.different_first_page_header_footer=False
doc.settings.odd_and_even_pages_header_footer=False

# Named reference override: monochrome academic typography, unshaded compact
# six-column results table, indented prose, and separate figure page.
def make_style(name,size,bold=False,after=0,line=1):
    st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Times New Roman'
    st.font.size=Pt(size)
    st.font.bold=bold
    st.font.italic=False
    st.font.color.rgb=RGBColor(0,0,0)
    for key in list(st.element.get_or_add_rPr().rFonts.attrib):
        if key.endswith('Theme'):del st.element.rPr.rFonts.attrib[key]
    st.paragraph_format.space_before=Pt(0)
    st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.line_spacing=line
    return st

normal=make_style('Normal',12,after=8,line=1.8)
normal.paragraph_format.first_line_indent=Inches(.5)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
make_style('Academic caption',11,after=5)
make_style('Academic table',11)
make_style('Academic note',10,after=8,line=1.15)
make_style('Institution',10,True,line=1)
make_style('Page number',9,line=1)

# Header text and floating school seal reproduce the reference masthead.
hp=sec.header.paragraphs[0]
hp.style=doc.styles['Institution']
hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
hp.paragraph_format.first_line_indent=Inches(0)
picture=hp.add_run().add_picture(str(ROOT/'tmp/bfp-survey/image1.png'),width=Inches(.56))
inline=picture._inline
blip=inline.xpath('.//a:blip')[0]
blip.append(OxmlElement('a:grayscl'))
inline.docPr.set('descr','St. Anthony\u2019s College seal displayed in grayscale')
anchor=OxmlElement('wp:anchor')
for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'0','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():anchor.set(k,v)
pos=OxmlElement('wp:simplePos');pos.set('x','0');pos.set('y','0');anchor.append(pos)
for axis,reference,offset in [('H','column',0),('V','paragraph',0)]:
    pos=OxmlElement(f'wp:position{axis}');pos.set('relativeFrom',reference)
    off=OxmlElement('wp:posOffset');off.text=str(offset);pos.append(off);anchor.append(pos)
anchor.append(inline.extent)
effect=OxmlElement('wp:effectExtent')
for k in ['l','t','r','b']:effect.set(k,'0')
anchor.append(effect)
anchor.append(OxmlElement('wp:wrapNone'))
for child in list(inline):anchor.append(child)
inline.getparent().replace(inline,anchor)
hp.add_run('St. Anthony\u2019s College')
for txt in ['INFORMATION TECHNOLOGY DEPARTMENT','San Jose, Antique']:
    p=sec.header.add_paragraph(txt,'Institution');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Inches(0)
pagep=sec.header.add_paragraph('','Page number')
pagep.paragraph_format.first_line_indent=Inches(0)
pagep.paragraph_format.line_spacing=1
pagep.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');pagep._p.append(fld)

def p(text,style='Normal',align=None):
    pp=doc.add_paragraph(text,style)
    if style!='Normal':pp.paragraph_format.first_line_indent=Inches(0)
    if align is not None:pp.alignment=align
    return pp

p('The following mock tabulation illustrates how survey responses from residents and BFP personnel may be presented for the proposed provincial fire response system. All values are simulated and are provided for document formatting only.')
cp=p('Table 1. Resident and BFP Personnel Mock Survey Results','Academic caption',WD_ALIGN_PARAGRAPH.CENTER)
cp.paragraph_format.keep_with_next=True
table=doc.add_table(rows=14,cols=6)
widths=[1600,1360,1360,1600,1360,1360]
apply_table_geometry(table,widths,cell_margins_dxa={'top':30,'bottom':30,'start':80,'end':80})
table.cell(0,0).merge(table.cell(1,0)).text='Residents\n(n = 30)'
table.cell(0,1).merge(table.cell(0,2)).text='Response'
table.cell(0,3).merge(table.cell(1,3)).text='BFP Personnel\n(n = 10)'
table.cell(0,4).merge(table.cell(0,5)).text='Response'
for j,v in [(1,'Yes'),(2,'No'),(4,'Yes'),(5,'No')]:table.cell(1,j).text=v
for i,(r,b) in enumerate(zip(resident,bfp),2):
    for j,val in enumerate([f'Q{i-1}',r,30-r,f'Q{i-1}',b,10-b]):table.cell(i,j).text=str(val)
for i,values in [(12,['Total',265,35,'Total',90,10]),(13,['Percentage','88.33%','11.67%','Percentage','90.00%','10.00%'])]:
    for j,v in enumerate(values):table.cell(i,j).text=str(v)
edges=OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
    e=OxmlElement(f'w:{edge}');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');edges.append(e)
table._tbl.tblPr.append(edges)
for i,row in enumerate(table.rows):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    if i<2:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    seen=set()
    for cell in row.cells:
        if cell._tc in seen:continue
        seen.add(cell._tc)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
        for cp in cell.paragraphs:
            cp.style=doc.styles['Academic table'];cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.first_line_indent=Inches(0)
            for r in cp.runs:r.bold=i<2 or i>=12

p('Note. Mock data only. Q1-Q10 refer to the ten proposed survey items.','Academic note')
p('Table 1 presents simulated responses from 30 residents and 10 BFP personnel. The questionnaire consists of ten items, each with a Yes or No response. Across 300 resident answers, 265 are Yes responses (88.33%) and 35 are No responses (11.67%). Across 100 BFP personnel answers, 90 are Yes responses (90.00%) and 10 are No responses (10.00%).')
p('The two groups provide a combined total of 400 item responses, comprising 355 Yes answers and 45 No answers. Percentages are based on the total number of answers across the ten questions, rather than the number of individual participants. These figures illustrate the tabulation format and do not represent actual survey findings.')

doc.add_page_break()
cp=doc.add_paragraph()
cp.paragraph_format.first_line_indent=Inches(0)
cp.paragraph_format.space_after=Pt(6)
cp.paragraph_format.line_spacing=1
cp.paragraph_format.keep_with_next=True
cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pic=cp.add_run().add_picture(str(TMP/'pie.png'),width=Inches(4.65))
pic._inline.docPr.set('descr','Monochrome pie chart of mock data: Yes, 355 responses or 88.75%; No, 45 responses or 11.25%.')
cp=p('Figure 1. Mock Survey Results in a Pie Chart','Academic caption',WD_ALIGN_PARAGRAPH.CENTER)
cp.runs[0].italic=True
cp.paragraph_format.space_after=Pt(14)
p('The pie chart in Figure 1 shows the overall distribution of the simulated responses from residents and BFP personnel. Of the 400 item responses, 88.75% are Yes answers and 11.25% are No answers. These percentages are calculated from the combined response counts in Table 1.')
p('The chart summarizes answers across the ten proposed survey items. It does not indicate the percentage of individual participants who support the system. Actual responses must replace these mock values before the table and figure are presented as research findings.')

doc.core_properties.title='Resident and BFP Personnel Mock Survey Results'
doc.core_properties.subject='Black-and-white combined survey table and pie chart; simulated data'
doc.save(OUT)
check=Document(OUT)
assert len(check.tables)==1 and len(check.tables[0].columns)==6
grid=[int(g.get(qn('w:w'))) for g in check.tables[0]._tbl.tblGrid]
for tr in check.tables[0]._tbl.tr_lst:
    idx=0
    for tc in tr.tc_lst:
        n=int(tc.tcPr.gridSpan.val) if tc.tcPr.gridSpan is not None else 1
        assert int(tc.tcPr.tcW.w)==sum(grid[idx:idx+n]);idx+=n
    assert idx==6
assert check.tables[0].cell(12,1).text=='265'
assert check.tables[0].cell(13,5).text=='10.00%'
print(f'Created {OUT}')
print('Verified six-column merged-header geometry and unchanged mock counts.')
